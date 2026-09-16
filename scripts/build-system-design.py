from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Hudi-Labs-System-Design.artifact.md"
OUTPUT = ROOT / "docs" / "Hudi-Labs-System-Design.generated.docx"


def main() -> None:
    document = Document()
    document.core_properties.title = "Hudi Labs Website System Design"
    document.core_properties.subject = "Static routes, content, accessibility and delivery"
    document.core_properties.author = "Hudi Labs / Coletivo Inspira"

    document.add_heading("Hudi Labs — System Design", level=1)
    document.add_paragraph(f"Last updated: {date.today():%B %d, %Y}")

    for line in SOURCE.read_text().splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
        else:
            document.add_paragraph("")

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
